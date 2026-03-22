"""
Roof Plus 導入評価書 自動生成スクリプト
経営者向け 中立的解説資料（Word形式）
"""

import io
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

# ==============================
# 会社ごとに変更するパラメータ
# ==============================
PARAMS = {
    "company_name": "株式会社〇〇",           # 会社名
    "industry": "製造業",                      # 業種
    "roof_area_m2": 500,                        # 屋根面積（㎡）
    "system_kw": 80,                            # 設置容量（kW）
    "annual_kwh": 88000,                        # 年間発電量（kWh）
    "current_unit_price": 25.0,                 # 現在の電力単価（円/kWh）
    "price_rise_optimistic": 0.01,              # 楽観シナリオ 上昇率/年
    "price_rise_base": 0.03,                    # 基準シナリオ 上昇率/年
    "price_rise_pessimistic": 0.05,             # 悲観シナリオ 上昇率/年
    "years": 20,                                # 分析期間（年）
    # --- 信用ランクごとの割賦単価（円/kWh）---
    "rank_rates": {
        "Aランク": 18.0,
        "Bランク": 20.0,
        "Cランク": 22.5,
    },
    "fixed_asset_tax_annual": 150000,           # 年間固定資産税（円）※需要家負担分
    "other_expense_annual": 50000,              # 年間その他経費（円）
    "credit_rank": "Bランク",                   # 対象会社の信用ランク
}

import matplotlib.font_manager as fm
_ipa_fonts = [f.fname for f in fm.fontManager.ttflist if 'IPAGothic' in f.name or 'ipag' in f.fname.lower()]
if _ipa_fonts:
    fm.fontManager.addfont(_ipa_fonts[0])
    _prop = fm.FontProperties(fname=_ipa_fonts[0])
    plt.rcParams['font.family'] = _prop.get_name()
else:
    plt.rcParams['font.family'] = 'sans-serif'

def yen(val):
    return f"¥{val:,.0f}"

def make_chart_cost_comparison(params):
    """電力コスト vs Roof Plus コスト 20年推移グラフ"""
    years = list(range(1, params["years"] + 1))
    unit0 = params["current_unit_price"]
    kwh = params["annual_kwh"]
    rank_rate = params["rank_rates"][params["credit_rank"]]
    fixed = params["fixed_asset_tax_annual"]
    other = params["other_expense_annual"]

    scenarios = {
        f"電力単価 楽観({int(params['price_rise_optimistic']*100)}%/年)": [],
        f"電力単価 基準({int(params['price_rise_base']*100)}%/年)": [],
        f"電力単価 悲観({int(params['price_rise_pessimistic']*100)}%/年)": [],
        f"Roof Plus合計({params['credit_rank']})": [],
    }

    for y in years:
        for rate, key in [
            (params["price_rise_optimistic"], f"電力単価 楽観({int(params['price_rise_optimistic']*100)}%/年)"),
            (params["price_rise_base"],       f"電力単価 基準({int(params['price_rise_base']*100)}%/年)"),
            (params["price_rise_pessimistic"],f"電力単価 悲観({int(params['price_rise_pessimistic']*100)}%/年)"),
        ]:
            price = unit0 * ((1 + rate) ** y)
            scenarios[key].append(price * kwh / 10000)  # 万円

        roof_total = (rank_rate * kwh + fixed + other) / 10000
        scenarios[f"Roof Plus合計({params['credit_rank']})"].append(roof_total)

    fig, ax = plt.subplots(figsize=(9, 5))
    colors = ["#4caf50", "#2196f3", "#f44336", "#ff9800"]
    styles = ["--", "-.", ":", "-"]
    widths = [1.5, 1.5, 1.5, 2.5]

    for (label, values), color, ls, lw in zip(scenarios.items(), colors, styles, widths):
        ax.plot(years, values, label=label, color=color, linestyle=ls, linewidth=lw)

    ax.set_xlabel("経過年数（年）", fontsize=10)
    ax.set_ylabel("年間コスト（万円）", fontsize=10)
    ax.set_title("年間コスト比較：電力購入 vs Roof Plus", fontsize=12, fontweight='bold')
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda x, _: f"{x:,.0f}万円"))
    ax.legend(fontsize=9, loc="upper left")
    ax.grid(axis='y', alpha=0.3)
    ax.set_xlim(1, params["years"])
    plt.tight_layout()

    buf = io.BytesIO()
    fig.savefig(buf, format='png', dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf

def make_chart_cumulative_savings(params):
    """累積削減効果グラフ"""
    years = list(range(1, params["years"] + 1))
    unit0 = params["current_unit_price"]
    kwh = params["annual_kwh"]
    rank_rate = params["rank_rates"][params["credit_rank"]]
    fixed = params["fixed_asset_tax_annual"]
    other = params["other_expense_annual"]
    roof_annual = rank_rate * kwh + fixed + other

    fig, ax = plt.subplots(figsize=(9, 5))
    colors = {"楽観": "#4caf50", "基準": "#2196f3", "悲観": "#f44336"}

    for label, rate in [
        ("楽観", params["price_rise_optimistic"]),
        ("基準", params["price_rise_base"]),
        ("悲観", params["price_rise_pessimistic"]),
    ]:
        cumulative = []
        total = 0
        for y in years:
            elec = unit0 * ((1 + rate) ** y) * kwh
            total += (elec - roof_annual)
            cumulative.append(total / 10000)
        ax.plot(years, cumulative, label=f"{label}シナリオ", color=colors[label], linewidth=2)

    ax.axhline(0, color='gray', linewidth=0.8, linestyle='--')
    ax.fill_between(years, 0, 0, alpha=0.1)
    ax.set_xlabel("経過年数（年）", fontsize=10)
    ax.set_ylabel("累積削減効果（万円）", fontsize=10)
    ax.set_title("累積コスト削減効果（電力購入との差額）", fontsize=12, fontweight='bold')
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(lambda x, _: f"{x:+,.0f}万円"))
    ax.legend(fontsize=9)
    ax.grid(axis='y', alpha=0.3)
    ax.set_xlim(1, params["years"])
    plt.tight_layout()

    buf = io.BytesIO()
    fig.savefig(buf, format='png', dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf

def make_chart_cost_breakdown(params):
    """Roof Plusコスト内訳 積み上げ棒グラフ"""
    kwh = params["annual_kwh"]
    rank = params["credit_rank"]
    rate = params["rank_rates"][rank]

    installment = rate * kwh / 10000
    tax = params["fixed_asset_tax_annual"] / 10000
    other = params["other_expense_annual"] / 10000

    labels = list(params["rank_rates"].keys())
    installments = [r * kwh / 10000 for r in params["rank_rates"].values()]
    taxes = [tax] * len(labels)
    others = [other] * len(labels)

    fig, ax = plt.subplots(figsize=(7, 4.5))
    x = np.arange(len(labels))
    w = 0.5

    b1 = ax.bar(x, installments, w, label="割賦代金", color="#2196f3")
    b2 = ax.bar(x, taxes, w, bottom=installments, label="固定資産税", color="#ff9800")
    b3 = ax.bar(x, others, w,
                bottom=[i + t for i, t in zip(installments, taxes)],
                label="その他経費", color="#9e9e9e")

    totals = [i + tax + other for i in installments]
    for xi, total in zip(x, totals):
        ax.text(xi, total + 0.5, f"{total:.1f}万円", ha='center', fontsize=9, fontweight='bold')

    ax.set_xticks(x)
    ax.set_xticklabels(labels, fontsize=10)
    ax.set_ylabel("年間コスト（万円）", fontsize=10)
    ax.set_title("Roof Plus 年間コスト内訳（信用ランク別）", fontsize=12, fontweight='bold')
    ax.legend(fontsize=9, loc="upper right")
    ax.grid(axis='y', alpha=0.3)

    # 対象ランクをハイライト
    target_idx = labels.index(rank)
    ax.get_xticklabels()[target_idx].set_color("#f44336")
    ax.get_xticklabels()[target_idx].set_fontweight('bold')

    plt.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format='png', dpi=150)
    plt.close(fig)
    buf.seek(0)
    return buf

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p

def add_note(doc, text):
    p = doc.add_paragraph(text)
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    return p

def generate_report(params, output_path="RoofPlus_評価書.docx"):
    doc = Document()

    # --- 余白設定 ---
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)

    rank = params["credit_rank"]
    rank_rate = params["rank_rates"][rank]
    kwh = params["annual_kwh"]
    fixed = params["fixed_asset_tax_annual"]
    other = params["other_expense_annual"]
    roof_annual = rank_rate * kwh + fixed + other

    # ==================
    # P1: 表紙
    # ==================
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Roof Plus 導入評価書")
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1a, 0x73, 0xe8)

    doc.add_paragraph()
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"対象企業：{params['company_name']}　様").font.size = Pt(14)

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run(
        f"業種：{params['industry']}　｜　設置容量：{params['system_kw']}kW　｜　信用ランク：{rank}\n"
        f"作成日：{datetime.date.today().strftime('%Y年%m月%d日')}"
    ).font.size = Pt(11)

    doc.add_paragraph()
    note = doc.add_paragraph(
        "※ 本資料は導入検討のための参考資料です。実際のコストは現地調査・審査結果により異なります。\n"
        "　 将来の電力単価は予測値であり、保証するものではありません。"
    )
    note.runs[0].font.size = Pt(9)
    note.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    # ==================
    # P2: スキーム概要
    # ==================
    add_heading(doc, "1｜Roof Plus スキームの概要", level=1)

    doc.add_paragraph(
        "Roof Plus は、あいおいニッセイ同和損保と共同開発した自家消費型太陽光発電の普及スキームです。\n"
        "需要家（導入企業）は初期費用ゼロで太陽光設備を屋根上に設置し、\n"
        "割賦払いにより毎月固定コストで自家消費電力を確保できます。"
    )

    table = doc.add_table(rows=5, cols=2)
    table.style = 'Table Grid'
    headers = ["項目", "内容"]
    rows_data = [
        ["スキーム名", "Roof Plus（ルーフプラス）"],
        ["初期費用", "不要（割賦払いで対応）"],
        ["設備所有", "割賦期間中は事業者所有・期間終了後に移転"],
        ["保険", "あいおいニッセイ同和損保による包括補償"],
        ["対象", "屋根面積が一定以上の事業者"],
    ]
    for i, (k, v) in enumerate(rows_data):
        row = table.rows[i]
        row.cells[0].text = k
        row.cells[1].text = v
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        set_cell_bg(row.cells[0], "EBF3FD")

    doc.add_paragraph()
    add_note(doc, "※ 詳細条件は契約書・重要事項説明書をご確認ください。")
    doc.add_page_break()

    # ==================
    # P3: コスト構造（内訳グラフ）
    # ==================
    add_heading(doc, "2｜Roof Plus コスト構造（信用ランク別）", level=1)

    doc.add_paragraph(
        "Roof Plusの年間コストは ①割賦代金 ②固定資産税 ③その他経費 の3要素で構成されます。\n"
        "割賦代金は需要家の信用ランクにより異なります（下グラフ）。"
    )

    chart1_buf = make_chart_cost_breakdown(params)
    doc.add_picture(chart1_buf, width=Inches(5.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    table2 = doc.add_table(rows=5, cols=3)
    table2.style = 'Table Grid'
    hdr = table2.rows[0]
    for cell, txt in zip(hdr.cells, ["費用項目", "金額（年間）", "備考"]):
        cell.text = txt
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_bg(cell, "1A73E8")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    breakdown_rows = [
        [f"割賦代金（{rank}）", yen(rank_rate * kwh), f"単価 {rank_rate}円/kWh × {kwh:,}kWh"],
        ["固定資産税", yen(fixed), "需要家負担分（概算）"],
        ["その他経費", yen(other), "保守・管理費等"],
        ["合計", yen(roof_annual), ""],
    ]
    for i, row_data in enumerate(breakdown_rows):
        row = table2.rows[i + 1]
        for j, val in enumerate(row_data):
            row.cells[j].text = val
        if i == 3:
            for cell in row.cells:
                cell.paragraphs[0].runs[0].font.bold = True
            set_cell_bg(row.cells[0], "FFF3E0")
            set_cell_bg(row.cells[1], "FFF3E0")
            set_cell_bg(row.cells[2], "FFF3E0")

    doc.add_page_break()

    # ==================
    # P4: 電力コスト比較グラフ
    # ==================
    add_heading(doc, "3｜年間コスト比較：電力購入 vs Roof Plus（20年推移）", level=1)

    doc.add_paragraph(
        "電力単価は近年上昇傾向にあります。以下では楽観・基準・悲観の3シナリオで\n"
        "「そのまま電力購入を続けた場合」と「Roof Plusを導入した場合」の年間コストを比較します。"
    )

    chart2_buf = make_chart_cost_comparison(params)
    doc.add_picture(chart2_buf, width=Inches(6.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    table3 = doc.add_table(rows=2, cols=4)
    table3.style = 'Table Grid'
    hdr3 = table3.rows[0]
    for cell, txt in zip(hdr3.cells, ["シナリオ", "上昇率/年", "20年目の単価（推計）", "20年目の年間電力費"]):
        cell.text = txt
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_bg(cell, "1A73E8")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    u0 = params["current_unit_price"]
    for i, (label, rate) in enumerate([
        ("楽観", params["price_rise_optimistic"]),
        ("基準", params["price_rise_base"]),
        ("悲観", params["price_rise_pessimistic"]),
    ]):
        price_20 = u0 * ((1 + rate) ** 20)
        cost_20 = price_20 * kwh
        row = table3.add_row()
        row.cells[0].text = f"{label}シナリオ"
        row.cells[1].text = f"{rate*100:.0f}%"
        row.cells[2].text = f"{price_20:.1f}円/kWh"
        row.cells[3].text = yen(cost_20)

    add_note(doc, f"※ 現在の電力単価：{u0}円/kWh　年間使用量（自家消費分）：{kwh:,}kWh として試算")
    doc.add_page_break()

    # ==================
    # P5: 累積削減効果
    # ==================
    add_heading(doc, "4｜累積コスト削減効果", level=1)

    doc.add_paragraph(
        "Roof Plus導入により削減できる累積コストを示します。\n"
        "グラフがプラスの領域に入ると「電力購入より安くなった」ことを意味します。"
    )

    chart3_buf = make_chart_cumulative_savings(params)
    doc.add_picture(chart3_buf, width=Inches(6.0))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 20年後の削減額サマリ
    doc.add_paragraph()
    table4 = doc.add_table(rows=4, cols=3)
    table4.style = 'Table Grid'
    hdr4 = table4.rows[0]
    for cell, txt in zip(hdr4.cells, ["シナリオ", "20年累積削減額", "判定"]):
        cell.text = txt
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_bg(cell, "1A73E8")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for i, (label, rate) in enumerate([
        ("楽観", params["price_rise_optimistic"]),
        ("基準", params["price_rise_base"]),
        ("悲観", params["price_rise_pessimistic"]),
    ]):
        total = sum(
            u0 * ((1 + rate) ** y) * kwh - roof_annual
            for y in range(1, 21)
        )
        row = table4.rows[i + 1]
        row.cells[0].text = f"{label}シナリオ"
        row.cells[1].text = yen(total)
        row.cells[2].text = "削減効果あり ✓" if total > 0 else "慎重に検討"
        if total > 0:
            set_cell_bg(row.cells[2], "E8F5E9")
        else:
            set_cell_bg(row.cells[2], "FFEBEE")

    doc.add_page_break()

    # ==================
    # P6: 経営者向けまとめ・チェックリスト
    # ==================
    add_heading(doc, "5｜経営者向け 導入判断チェックリスト", level=1)

    checks = [
        ("メリット", [
            "初期投資ゼロで太陽光発電を導入できる",
            "電力コストの一部を固定化でき、価格変動リスクを低減できる",
            "CO₂削減・ESG対応・脱炭素への取り組みをアピールできる",
            "固定資産税の優遇措置が適用される可能性がある（要確認）",
        ]),
        ("注意点・リスク", [
            "割賦代金は信用ランクにより異なる（審査が必要）",
            "電力単価が低下した場合は相対的なメリットが減少する",
            "屋根の構造・強度・方位により設置できない場合がある",
            "割賦期間中の設備移設・解約には費用が発生する可能性がある",
        ]),
        ("導入前に確認すること", [
            "現在の年間電力使用量と電気代の確認",
            "屋根面積・構造の現地調査の実施",
            "信用審査の申込み・ランク確認",
            "契約内容・重要事項説明書の精読",
        ]),
    ]

    for category, items in checks:
        p = doc.add_paragraph()
        run = p.add_run(f"■ {category}")
        run.font.bold = True
        run.font.size = Pt(11)
        if category == "メリット":
            run.font.color.rgb = RGBColor(0x1a, 0x73, 0xe8)
        elif category == "注意点・リスク":
            run.font.color.rgb = RGBColor(0xf4, 0x43, 0x36)
        else:
            run.font.color.rgb = RGBColor(0xff, 0x99, 0x00)

        for item in items:
            bullet = doc.add_paragraph(style='List Bullet')
            bullet.add_run(item).font.size = Pt(10)

    doc.add_paragraph()
    add_note(doc, "※ 本資料の数値は試算値です。導入にあたっては専門家・担当者にご相談ください。")
    add_note(doc, f"※ 作成：一般社団法人 日本再生可能エネルギー地域資源開発機構　/ {datetime.date.today().strftime('%Y年%m月%d日')}")

    doc.save(output_path)
    print(f"✅ 評価書を出力しました：{output_path}")
    return output_path

if __name__ == "__main__":
    generate_report(PARAMS, output_path="RoofPlus_評価書_サンプル.docx")
